import json
import tempfile
import uuid
import os
from pathlib import Path
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel

try:
    import openpyxl
except ImportError:
    openpyxl = None

router = APIRouter(prefix="/excel")

CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

class CreateExcelRequest(BaseModel):
    title: str | None = None
    content: str

@router.post("/replace")
async def replace_excel(
    file: UploadFile = File(...),
    replacements: str = Form(...),
):
    try:
        replacement_items = json.loads(replacements)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail=f"Invalid replacements JSON: {exc}") from exc

    if not isinstance(replacement_items, list):
        raise HTTPException(status_code=400, detail="replacements must be a JSON array")

    if not openpyxl:
        raise HTTPException(status_code=500, detail="openpyxl library is not installed.")

    temp_path = Path(tempfile.mkdtemp(prefix="ang-excel-"))
    try:
        original_name = Path(file.filename or "document.xlsx").name
        input_path = temp_path / f"{uuid.uuid4()}-{original_name}"
        output_path = temp_path / f"{input_path.stem}-edited.xlsx"

        input_path.write_bytes(await file.read())

        wb = openpyxl.load_workbook(filename=str(input_path))
        for rep in replacement_items:
            find_text = rep.get("find")
            replace_text = rep.get("replace")
            if not find_text or replace_text is None:
                continue

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str):
                            if find_text in cell.value:
                                cell.value = cell.value.replace(find_text, str(replace_text))

        wb.save(str(output_path))
        wb.close()

        download_name = f"{Path(original_name).stem}-edited.xlsx"
        return _temp_file_response(
            output_path,
            temp_path,
            media_type=CONTENT_TYPES["xlsx"],
            filename=download_name,
        )
    except Exception:
        _cleanup_temp_dir(temp_path)
        raise

@router.post("/create")
async def create_excel(req: CreateExcelRequest):
    content = (req.content or "").strip()
    if not content:
        raise HTTPException(status_code=400, detail="content is required")

    if not openpyxl:
        raise HTTPException(status_code=500, detail="openpyxl library is not installed.")

    temp_path = Path(tempfile.mkdtemp(prefix="ang-excel-"))
    try:
        base_name = req.title.strip() if req.title and req.title.strip() else f"Created_{uuid.uuid4().hex[:8]}"
        if base_name.lower().endswith(".xlsx"):
            base_name = base_name[:-5]

        output_path = temp_path / f"{base_name}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        
        for row_idx, line in enumerate(content.split("\n"), start=1):
            for col_idx, cell_value in enumerate(line.split("\t"), start=1):
                ws.cell(row=row_idx, column=col_idx, value=cell_value)
                
        wb.save(str(output_path))
        wb.close()

        download_name = f"{base_name}.xlsx"
        return _temp_file_response(
            output_path,
            temp_path,
            media_type=CONTENT_TYPES["xlsx"],
            filename=download_name,
        )
    except Exception:
        _cleanup_temp_dir(temp_path)
        raise

def _temp_file_response(file_path: Path, temp_dir: Path, **kwargs):
    from starlette.background import BackgroundTask
    import shutil
    return FileResponse(
        path=file_path,
        background=BackgroundTask(_cleanup_temp_dir, temp_dir),
        **kwargs,
    )

def _cleanup_temp_dir(temp_dir: Path):
    import shutil
    try:
        if temp_dir.exists():
            shutil.rmtree(temp_dir, ignore_errors=True)
    except Exception as exc:
        print(f"Failed to clean up temp dir {temp_dir}: {exc}")
