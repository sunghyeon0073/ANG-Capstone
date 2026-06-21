import json
import tempfile
import uuid
import os
from pathlib import Path
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from typing import Dict, Any
from template_generator import generate_document_from_template, extract_placeholders

try:
    import docx
except ImportError:
    docx = None

router = APIRouter(prefix="/docx")

CONTENT_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

class CreateDocxRequest(BaseModel):
    title: str | None = None
    content: str

class GenerateTemplateRequest(BaseModel):
    template_name: str
    form_data: Dict[str, Any]
    title: str | None = None

@router.post("/parse-template")
async def parse_template_docx(file: UploadFile = File(...)):
    if not docx:
        raise HTTPException(status_code=500, detail="python-docx library is not installed.")
    
    try:
        contents = await file.read()
        placeholders = extract_placeholders(contents)
        return {"fields": [{"id": p, "label": p, "type": "text"} for p in placeholders]}
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Failed to parse template: {e}")

@router.post("/replace")
async def replace_docx(
    file: UploadFile = File(...),
    replacements: str = Form(...),
):
    try:
        replacement_items = json.loads(replacements)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail=f"Invalid replacements JSON: {exc}") from exc

    if not isinstance(replacement_items, list):
        raise HTTPException(status_code=400, detail="replacements must be a JSON array")

    if not docx:
        raise HTTPException(status_code=500, detail="python-docx library is not installed.")

    temp_path = Path(tempfile.mkdtemp(prefix="ang-docx-"))
    try:
        original_name = Path(file.filename or "document.docx").name
        input_path = temp_path / f"{uuid.uuid4()}-{original_name}"
        output_path = temp_path / f"{input_path.stem}-edited.docx"

        input_path.write_bytes(await file.read())

        doc = docx.Document(str(input_path))
        for rep in replacement_items:
            find_text = rep.get("find")
            replace_text = rep.get("replace")
            if not find_text or replace_text is None:
                continue

            # Basic paragraph replacement
            for paragraph in doc.paragraphs:
                if find_text in paragraph.text:
                    paragraph.text = paragraph.text.replace(find_text, replace_text)

            # Table replacement
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if find_text in cell.text:
                            cell.text = cell.text.replace(find_text, replace_text)

        doc.save(str(output_path))

        download_name = f"{Path(original_name).stem}-edited.docx"
        return _temp_file_response(
            output_path,
            temp_path,
            media_type=CONTENT_TYPES["docx"],
            filename=download_name,
        )
    except Exception:
        _cleanup_temp_dir(temp_path)
        raise

@router.post("/create")
async def create_docx(req: CreateDocxRequest):
    content = (req.content or "").strip()
    if not content:
        raise HTTPException(status_code=400, detail="content is required")

    if not docx:
        raise HTTPException(status_code=500, detail="python-docx library is not installed.")

    temp_path = Path(tempfile.mkdtemp(prefix="ang-docx-"))
    try:
        base_name = req.title.strip() if req.title and req.title.strip() else f"Created_{uuid.uuid4().hex[:8]}"
        if base_name.lower().endswith(".docx"):
            base_name = base_name[:-5]

        output_path = temp_path / f"{base_name}.docx"

        doc = docx.Document()
        for p in content.split("\n"):
            doc.add_paragraph(p)
            
        doc.save(str(output_path))

        download_name = f"{base_name}.docx"
        return _temp_file_response(
            output_path,
            temp_path,
            media_type=CONTENT_TYPES["docx"],
            filename=download_name,
        )
    except Exception:
        _cleanup_temp_dir(temp_path)
        raise

@router.post("/generate-template")
async def generate_template_docx(
    file: UploadFile = File(...),
    form_data: str = Form(...),
    title: str = Form(None)
):
    if not docx:
        raise HTTPException(status_code=500, detail="python-docx library is not installed.")
    
    try:
        import json
        form_data_dict = json.loads(form_data)
        template_bytes = await file.read()
        
        doc_stream = generate_document_from_template(template_bytes, form_data_dict)
        
        base_name = title.strip() if title and title.strip() else f"Document_{uuid.uuid4().hex[:8]}"
        if not base_name.lower().endswith(".docx"):
            base_name += ".docx"
            
        import urllib.parse
        encoded_name = urllib.parse.quote(base_name)
        return StreamingResponse(
            doc_stream,
            media_type=CONTENT_TYPES["docx"],
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_name}"}
        )
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Failed to generate template: {e}")

def _temp_file_response(file_path: Path, temp_dir: Path, **kwargs):
    from starlette.background import BackgroundTask
    import shutil
    return FileResponse(
        path=file_path,
        background=BackgroundTask(_cleanup_temp_dir, temp_dir),
        **kwargs,
    )

def _cleanup_temp_dir(temp_dir: Path):
    import shutil
    try:
        if temp_dir.exists():
            shutil.rmtree(temp_dir, ignore_errors=True)
    except Exception as exc:
        print(f"Failed to clean up temp dir {temp_dir}: {exc}")
