import os
from docx import Document

def create_meeting_minutes_template():
    doc = Document()
    doc.add_heading('회의록', 0)
    
    doc.add_heading('1. 회의 개요', level=1)
    doc.add_paragraph('일시: {{MEETING_DATE}}')
    doc.add_paragraph('장소: {{MEETING_PLACE}}')
    doc.add_paragraph('참석자: {{ATTENDEES}}')
    
    doc.add_heading('2. 주요 안건', level=1)
    doc.add_paragraph('{{AGENDA}}')
    
    doc.add_heading('3. 회의 내용', level=1)
    doc.add_paragraph('{{MEETING_NOTES}}')
    
    doc.add_heading('4. 향후 계획 및 결정 사항', level=1)
    doc.add_paragraph('{{DECISIONS}}')
    
    doc.save('templates/meeting_minutes_template.docx')

def create_business_plan_template():
    doc = Document()
    doc.add_heading('사업계획서', 0)
    
    doc.add_heading('1. 프로젝트 개요', level=1)
    doc.add_paragraph('프로젝트명: {{PROJECT_NAME}}')
    doc.add_paragraph('작성자: {{AUTHOR}}')
    
    doc.add_heading('2. 추진 배경 및 목적', level=1)
    doc.add_paragraph('{{BACKGROUND}}')
    
    doc.add_heading('3. 사업 내용', level=1)
    doc.add_paragraph('{{CONTENT}}')
    
    doc.add_heading('4. 예산 및 일정', level=1)
    doc.add_paragraph('{{BUDGET_AND_SCHEDULE}}')
    
    doc.add_heading('5. 기대 효과', level=1)
    doc.add_paragraph('{{EXPECTED_EFFECTS}}')
    
    doc.save('templates/business_plan_template.docx')

if __name__ == "__main__":
    if not os.path.exists('templates'):
        os.makedirs('templates')
    create_meeting_minutes_template()
    create_business_plan_template()
    print("Templates created successfully.")
