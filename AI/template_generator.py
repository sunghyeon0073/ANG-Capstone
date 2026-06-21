import os
import io
import re
from docx import Document

def replace_text_in_paragraphs(paragraphs, data):
    for p in paragraphs:
        for key, value in data.items():
            placeholder = f"{{{{{key}}}}}"
            if placeholder in p.text:
                inline = p.runs
                for i in range(len(inline)):
                    if placeholder in inline[i].text:
                        inline[i].text = inline[i].text.replace(placeholder, str(value))
                
                # Sometime docx splits a single placeholder across multiple runs. 
                # If the simple replace above didn't work because of splitting, 
                # we just replace the entire paragraph text, which destroys some internal formatting but works reliably.
                if placeholder in p.text:
                    p.text = p.text.replace(placeholder, str(value))

def extract_placeholders(file_bytes: bytes):
    """
    Reads a docx byte stream and extracts all unique placeholders formatted as {{KEY}}.
    """
    doc = Document(io.BytesIO(file_bytes))
    placeholders = set()
    
    def find_in_paragraphs(paragraphs):
        for p in paragraphs:
            matches = re.findall(r'\{\{(.*?)\}\}', p.text)
            for m in matches:
                placeholders.add(m.strip())

    find_in_paragraphs(doc.paragraphs)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                find_in_paragraphs(cell.paragraphs)
                
    return sorted(list(placeholders))

def generate_document_from_template(template_bytes: bytes, form_data):
    """
    Reads a template docx byte stream, replaces placeholders like {{KEY}} with form_data[KEY],
    and returns the byte stream of the generated docx.
    """
    doc = Document(io.BytesIO(template_bytes))
    
    replace_text_in_paragraphs(doc.paragraphs, form_data)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_text_in_paragraphs(cell.paragraphs, form_data)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
